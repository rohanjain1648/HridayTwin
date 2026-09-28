"""Builds docs/HridayTwin_Presentation.pptx from live metrics.   python docs/make_deck.py"""

import json
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
DOCS, ART = ROOT / "docs", ROOT / "artifacts"
M = json.loads((ART / "metrics.json").read_text())

PAPER, INK, INK2, SLATE = RGBColor(0xF2, 0xF7, 0xF5), RGBColor(0x15, 0x30, 0x4B), RGBColor(0x3A, 0x50, 0x68), RGBColor(0x6B, 0x7C, 0x8F)
TEAL, TEAL_W, RED, AMBER, WHITE = RGBColor(0x00, 0x9E, 0x9A), RGBColor(0xDF, 0xF2, 0xF0), RGBColor(0xC8, 0x41, 0x2F), RGBColor(0xD9, 0x91, 0x2B), RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Segoe UI"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]


def slide(title=None, kicker=None):
    s = prs.slides.add_slide(BLANK)
    bg = s.background.fill
    bg.solid()
    bg.fore_color.rgb = PAPER
    if title:
        text(s, title, 0.6, 0.45, 12.1, 0.9, 30, INK, bold=True)
    if kicker:
        text(s, kicker, 0.6, 1.25, 12.1, 0.6, 15, SLATE)
    return s


def text(s, t, x, y, w, h, size=16, color=INK, bold=False, align=PP_ALIGN.LEFT):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    lines = t if isinstance(t, list) else [t]
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = ln
        r.font.size, r.font.bold, r.font.color.rgb, r.font.name = Pt(size), bold, color, FONT
        p.space_after = Pt(6)
    return tb


def bullets(s, items, x, y, w, h, size=16, color=INK2):
    return text(s, [f"•  {i}" for i in items], x, y, w, h, size, color)


def rect(s, x, y, w, h, fill=WHITE, line=None):
    r = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    r.adjustments[0] = 0.08
    r.fill.solid()
    r.fill.fore_color.rgb = fill
    if line is None:
        r.line.fill.background()
    else:
        r.line.color.rgb = line
    r.shadow.inherit = False
    return r


def stat(s, value, label, x, y, w=2.9, color=INK):
    rect(s, x, y, w, 1.45)
    text(s, value, x + 0.2, y + 0.12, w - 0.4, 0.7, 32, color, bold=True)
    text(s, label, x + 0.2, y + 0.82, w - 0.4, 0.6, 12.5, SLATE)


def image(s, path, x, y, w=None, h=None):
    kw = {}
    if w:
        kw["width"] = Inches(w)
    if h:
        kw["height"] = Inches(h)
    return s.shapes.add_picture(str(path), Inches(x), Inches(y), **kw)


def notes(s, t):
    s.notes_slide.notes_text_frame.text = t


t, lt, tv, se, ab = M["test"], M["lead_time_medium"], M["twin_validation"], M["sentinel"], M["ablation"]

# Example patient for the what-if slide
from twin_engine.service import TwinService  # noqa: E402
svc = TwinService()
roster = svc.roster()
example = None
for r in roster:
    if r["tier"] in ("high", "medium"):
        tw = svc.twin(r["id"])
        best = tw["recommendations"][0]
        if example is None or best["risk_reduction"] > example[2]["risk_reduction"]:
            example = (r, tw, best)
er, etw, ebest = example
combo = svc.simulate(er["id"], {"breathing_min": 20, "sleep_extra_h": 1, "adherence": True, "low_salt": False, "steps_extra": 0})

# 1 ─ Title
s = slide()
rect(s, 0.6, 0.6, 0.75, 0.75, INK)
text(s, "♥", 0.6, 0.62, 0.75, 0.7, 26, WHITE, bold=True, align=PP_ALIGN.CENTER)
text(s, "HridayTwin", 0.6, 1.9, 12, 1.3, 66, INK, bold=True)
text(s, "A stress-aware digital twin that predicts hypertensive surges a day ahead, and shows the doctor how to prevent them",
     0.6, 3.25, 10.5, 1.3, 24, INK2)
text(s, ["Team [Team name]  |  [College / incubator]", "Built on MyZenn (myzenn.com)",
         "Happiest Health Digital Twin Challenge 2026, hypertension track"], 0.6, 5.2, 11, 1.6, 16, SLATE)
notes(s, "Hriday means heart. We built a working digital twin for hypertension on top of MyZenn, our live mental-wellness platform.")

# 2 ─ Problem
s = slide("India's hypertension problem happens between clinic visits",
          "Blood pressure is checked every few months. Dangerous surges happen on ordinary days in between.")
stat(s, "188 M", "Indian adults aged 30–79 living with hypertension (WHO, 2023)", 0.6, 2.2, 3.9)
stat(s, "~15%", "have it under control (WHO, 2023)", 4.75, 2.2, 3.9, RED)
stat(s, "#1", "modifiable risk factor for stroke and heart attack", 8.9, 2.2, 3.85)
bullets(s, ["Clinic readings are snapshots; home cuffs are used irregularly and read in isolation",
            "Surges are driven by everyday triggers: stress, poor sleep, missed doses, salty meals",
            "Doctors see what happened last quarter, not what will happen tomorrow",
            "Care is reactive: the first signal is often an ER visit"], 0.6, 4.1, 12, 3, 18)
notes(s, "WHO Global Report on Hypertension 2023, India profile.")

# 3 ─ Insight
s = slide("The insight: surges announce themselves in mind and body first",
          "Wearables and a mental-wellness app see the triggers hours before the cuff sees the pressure.")
cols = [("Stress", "HRV falls and resting HR rises as stress builds over 1–2 days"),
        ("Sleep", "Short, broken nights accumulate sleep debt that raises next-day BP"),
        ("Adherence", "A missed-dose streak lets the drug level fall"),
        ("Salt", "Festival and restaurant meals load sodium for 24–36 h")]
for i, (h, d) in enumerate(cols):
    x = 0.6 + i * 3.1
    rect(s, x, 2.2, 2.9, 2.4)
    text(s, h, x + 0.2, 2.35, 2.5, 0.6, 22, TEAL, bold=True)
    text(s, d, x + 0.2, 3.0, 2.55, 1.6, 15, INK2)
rect(s, 0.6, 5.0, 12.1, 1.7, TEAL_W)
text(s, ["MyZenn already measures the mind layer: mood check-ins, stress, HRV over Bluetooth, and guided breathing,",
         "meditation and yoga tools that lower sympathetic tone. HridayTwin connects that layer to the cardiologist."],
     0.85, 5.2, 11.6, 1.4, 17, INK)

# 4 ─ Solution
s = slide("HridayTwin: a virtual copy of each patient's blood pressure",
          "Four things a doctor can do with it, all working in the prototype.")
caps = [("Predict", f"Probability of a sustained surge (≥160/100) in the next 24 h, updated every 6 h. AUROC {t['auroc']:.2f}."),
        ("Explain", "Which drivers push this patient's risk up today, plus their personal sensitivities in mmHg."),
        ("Simulate", "What-if sliders: more sleep, MyBreath minutes, adherence, low salt. The twin re-runs live."),
        ("Act", "One-click care plans pushed to the patient's MyZenn app; MySafe escalation in crisis range.")]
for i, (h, d) in enumerate(caps):
    x = 0.6 + i * 3.1
    rect(s, x, 2.2, 2.9, 3.6, WHITE, TEAL if i == 2 else None)
    text(s, h, x + 0.2, 2.4, 2.5, 0.6, 24, INK, bold=True)
    text(s, d, x + 0.2, 3.15, 2.55, 2.6, 15, INK2)
text(s, "Doctor sees the twin in a clinician console. The patient lives the plan in MyZenn.", 0.6, 6.3, 12, 0.6, 16, SLATE)

# 5 ─ Data
c = M["cohort"]
s = slide("Two data streams, fused. Synthetic by design.",
          "No real patient data is used (DPDP Act and HIPAA). A causal simulator generates physiologically coherent patients.")
rect(s, 0.6, 2.1, 5.9, 3.3)
text(s, "Static / historical EHR", 0.85, 2.25, 5.5, 0.6, 20, INK, bold=True)
bullets(s, ["Synthea-style tables: patients, conditions, observations, medications",
            "Diagnoses coded in SNOMED-CT; 3-point lab history in LOINC",
            "eGFR, HbA1c, lipids, UACR, Na/K, clinic BP, BMI",
            "Genetic markers: AGT M235T, ACE I/D, CYP11B2, GNB3, polygenic score"], 0.85, 2.85, 5.5, 2.5, 14)
rect(s, 6.8, 2.1, 5.9, 3.3)
text(s, "Dynamic wearable + IoT (hourly)", 7.05, 2.25, 5.5, 0.6, 20, INK, bold=True)
bullets(s, ["Smartwatch: heart rate, HRV (RMSSD), steps, stress score, sleep stages",
            "Home BP cuff, morning and evening (often skipped, as in real life)",
            "Smart pillbox dose log; MyZenn food log",
            "MyZenn mood check-ins and breathing / meditation sessions"], 7.05, 2.85, 5.5, 2.5, 14)
stat(s, f"{c['patients']}", "synthetic patients, 60 days each", 0.6, 5.7, 2.9)
stat(s, "576k", "hourly wearable rows", 3.7, 5.7, 2.9)
stat(s, f"{c['events_total']:,}", "ground-truth surge events", 6.8, 5.7, 2.9)
stat(s, f"{c['features']}", "fused features per prediction", 9.9, 5.7, 2.8)
notes(s, "Ground truth comes from a hidden ambulatory BP signal the model never sees. The model only gets what a real clinic would get.")

# 6 ─ Architecture
s = slide("Architecture")
image(s, DOCS / "HridayTwin_Architecture.png", 0.6, 1.25, w=12.1)

# 7 ─ How the twin works
s = slide("A hybrid twin: machine learning for accuracy, physiology for what-if",
          "Neither half alone is a digital twin. Together they predict and simulate.")
rect(s, 0.6, 2.1, 5.9, 4.6)
text(s, "Risk engine (ML)", 0.85, 2.25, 5.5, 0.6, 20, INK, bold=True)
bullets(s, ["71 strictly causal features every 6 h: uses only data before the prediction time",
            "LightGBM with 36 clinical monotonicity constraints (for example, more doses taken can never raise risk)",
            "Platt-calibrated probability; TreeSHAP explanations grouped into clinical drivers",
            "Patient-level train / validation / test split (240 / 60 / 100)"], 0.85, 2.85, 5.5, 3.8, 14)
rect(s, 6.8, 2.1, 5.9, 4.6, WHITE, TEAL)
text(s, "Mechanistic twin (physiology)", 7.05, 2.25, 5.5, 0.6, 20, TEAL, bold=True)
text(s, "SBP = θ₀ + θ_stress·stress + θ_sleep·sleep debt + θ_drug·drug level + θ_salt·salt load + θ_am·morning",
     7.05, 2.85, 5.5, 0.9, 13, INK)
bullets(s, ["Bayesian calibration on each patient's own cuff readings, shrunk to a population prior",
            "Personal sensitivities with uncertainty, and a phenotype",
            "What-if: the scenario changes physiology, the twin moves BP, the ML re-scores risk",
            "24-hour BP projection with an 80% band"], 7.05, 3.75, 5.5, 3, 14)

# 8 ─ Results
s = slide("Results on 100 patients the model never saw",
          f"{c['samples']:,} prediction points. Every number below is reproducible with one command.")
stat(s, f"{t['auroc']:.3f}", "AUROC, surge in the next 24 h", 0.6, 2.0)
stat(s, f"{round(lt['event_recall'] * 100)}%", "of surges flagged in advance (risk ≥ 30%)", 3.7, 2.0, color=TEAL)
stat(s, f"{lt['median_lead_h']:.0f} h", "median warning before onset", 6.8, 2.0, color=TEAL)
stat(s, f"{t['brier']:.3f}", "Brier score (well calibrated)", 9.9, 2.0, w=2.8)
image(s, ART / "plots" / "performance.png", 0.6, 3.75, w=12.1)
notes(s, f"At the 55% high tier: sensitivity {t['at_high']['sensitivity']:.2f}, specificity {t['at_high']['specificity']:.2f}, "
         f"PPV {t['at_high']['ppv']:.2f}, about {t['at_high']['alerts_per_patient_week']:.1f} alert-days per patient-week. "
         f"{round(lt['pct_caught_6h_plus'] * 100)}% of caught surges were flagged at least 6 hours ahead.")

# 9 ─ Fusion ablation
s = slide("Fusion matters: every stream adds signal",
          "Same model, same patients, different inputs.")
image(s, ART / "plots" / "ablation.png", 0.6, 2.0, w=7.6)
fus = ab["HridayTwin fusion (EHR + streams)"]["auroc"]
bullets(s, [f"Fusion {fus:.3f} vs home cuff alone {ab['Home cuff only (standard care)']['auroc']:.3f} vs today's rule of thumb {ab['Clinical rule (last home SBP)']['auroc']:.3f}",
            f"EHR alone ({ab['EHR only (static)']['auroc']:.2f}) knows who is at risk; streams know when",
            f"Wearables + MyZenn alone reach {ab['Wearables + MyZenn only']['auroc']:.2f} with no BP cuff at all: useful for the many patients who rarely measure"],
        8.5, 2.2, 4.3, 4.5, 15)

# 10 ─ Twin validation
s = slide("The twin learns each patient's real physiology",
          "Every simulated patient has hidden sensitivities. The twin sees only cuff readings and wearables.")
image(s, ART / "plots" / "twin_validation.png", 0.6, 1.95, h=4.9)
bullets(s, [f"Stress sensitivity recovered with Spearman ρ = {tv['corr_stress_sensitivity']:.2f}",
            f"Medication effect ρ = {tv['corr_drug_effect']:.2f}; salt ρ = {tv['corr_salt_sensitivity']:.2f}; sleep ρ = {tv['corr_sleep_sensitivity']:.2f}",
            f"True main driver is in the twin's top two for {round(tv['phenotype_top2_accuracy'] * 100)}% of patients",
            "This is what makes the what-if simulator personal rather than generic",
            "Honest limit: single-driver phenotype accuracy is close to the majority baseline; we report the top-two view"],
        6.6, 2.1, 6.2, 4.8, 15)

# 11 ─ Console
s = slide("The clinician console", "Triage the panel, open the virtual patient, understand, simulate, act.")
image(s, DOCS / "screenshots" / "console.png", 0.6, 1.8, w=8.4)
bullets(s, ["Roster ranked by 24 h surge risk with phenotype badges",
            "Blood pressure twin: cuff history flowing into the projected next 24 h",
            "Risk dial and today's snapshot",
            "Why this risk: driver bars, top signals, personal sensitivities",
            "Twin clock: replay the past two weeks and watch alerts fire before surges",
            "Live monitor: real-time stream, webcam heart rate, Bluetooth strap"],
        9.2, 1.9, 3.7, 5.2, 13.5)

# 11b ─ Real-time data
s = slide("Real-time data flows into the twin", "Three live sources, one twin. Every reading is stored; risk updates in seconds.")
image(s, DOCS / "screenshots" / "live_monitor.png", 0.6, 1.8, w=8.4)
bullets(s, ["Streaming replay over WebSocket: one hour of wearable, cuff and app data per tick; the twin re-scores every tick",
            "Alerts fire with drivers and the twin's recommended protocol, while the true BP (dotted, never seen by the model) is still below 160",
            "Webcam heart rate (rPPG, POS method) runs in the browser: no wearable, video never leaves the device; within 3 bpm in tests",
            "Bluetooth strap gives true beat-to-beat HRV (MyZenn MyPulse hook)",
            "All live readings persist in SQLite with source and signal quality"],
        9.2, 1.9, 3.7, 5.3, 13)

# 12 ─ What-if + MySentinel
s = slide("What-if, then act: the twin picks the intervention",
          f"Example: {er['name']} ({er['age']}{er['sex']}, {er['city']}), {etw['phenotype']['primary'].lower()}.")
rect(s, 0.6, 2.05, 6.0, 4.7)
text(s, "Simulated on this patient's twin", 0.85, 2.2, 5.6, 0.6, 18, INK, bold=True)
rows = [f"{r['name']} ({r['myzenn_tool']}): risk {round(etw['risk'] * 100)}% → {round(r['risk_after'] * 100)}%"
        if r["risk_reduction"] > 0.005 else f"{r['name']} ({r['myzenn_tool']}): no benefit predicted for this patient"
        for r in etw["recommendations"]]
bullets(s, rows, 0.85, 2.8, 5.6, 2.6, 14)
text(s, f"Combined (MyBreath 20 min + 1 h more sleep + every dose): {round(combo['risk_before'] * 100)}% → {round(combo['risk_after'] * 100)}%, "
        f"projected peak {combo['projection_before']['peak_sbp']:.0f} → {combo['projection_after']['peak_sbp']:.0f} mmHg",
     0.85, 5.55, 5.6, 1.1, 14, TEAL, bold=True)
rect(s, 6.9, 2.05, 5.8, 4.7, TEAL_W)
text(s, "MySentinel learns what works", 7.15, 2.2, 5.4, 0.6, 18, INK, bold=True)
bullets(s, ["LinUCB contextual bandit, ported from MyZenn's MyCompass",
            f"Pre-trained in silico on {se['train_rounds']:,} twin cases",
            f"Average risk cut: {se['mean_risk_reduction_bandit'] * 100:.0f} pts learned vs {se['mean_risk_reduction_random'] * 100:.0f} pts random "
            f"(best possible {se['mean_risk_reduction_oracle'] * 100:.0f})",
            "Policy tiers: crisis → MySafe call; high → clinician alert + protocol; medium → patient nudge",
            "Keeps learning from real outcomes after deployment"], 7.15, 2.8, 5.4, 3.9, 14)

# 13 ─ Trust
s = slide("Built to be trusted", "Privacy, safety and clinical sense are part of the design, not an afterthought.")
items = [("Privacy (DPDP)", "Synthetic training data. In production, MyZenn's edge model keeps raw journal and mood data on the phone; only derived scores sync."),
         ("Clinician in the loop", "The twin recommends; the doctor prescribes. Every action is logged."),
         ("Clinically constrained", "36 monotonicity rules and physiological sign constraints; tested automatically (no future leakage, adherence never raises risk)."),
         ("Safety net", "Crisis-range readings trigger MySafe escalation. Notifications never reveal health details on a lock screen."),
         ("Interoperable", "Synthea / FHIR-style EHR with SNOMED-CT and LOINC; ready for ABDM health records."),
         ("Honest", "Research prototype, not a medical device. Clinical validation is the next step.")]
for i, (h, d) in enumerate(items):
    x, y = 0.6 + (i % 3) * 4.1, 2.05 + (i // 3) * 2.45
    rect(s, x, y, 3.9, 2.25)
    text(s, h, x + 0.2, y + 0.15, 3.5, 0.5, 17, TEAL, bold=True)
    text(s, d, x + 0.2, y + 0.7, 3.55, 1.6, 13, INK2)

# 14 ─ Impact & roadmap
s = slide("From prototype to patients", "MyZenn is live today; HridayTwin is the bridge to cardiometabolic care.")
steps = [("Now", "Working prototype: data simulator, hybrid twin, clinician console, MyZenn hooks"),
         ("3 months", "Retrain on MIMIC-IV and open wearable datasets; pilot design with a Bengaluru hospital"),
         ("6 months", "Prospective pilot: 200 patients with ambulatory BP as ground truth; measure surges averted"),
         ("12 months", "CDSCO software-as-medical-device pathway; ABDM integration; insurer and clinic partnerships")]
for i, (h, d) in enumerate(steps):
    x = 0.6 + i * 3.1
    rect(s, x, 2.1, 2.9, 2.9, WHITE, TEAL if i == 0 else None)
    text(s, h, x + 0.2, 2.25, 2.5, 0.6, 22, INK, bold=True)
    text(s, d, x + 0.2, 2.9, 2.55, 2.1, 14, INK2)
bullets(s, ["Impact metric: surges predicted and prevented per 1,000 patient-months",
            "Model: B2B2C through clinics, corporate wellness and insurers; patients use MyZenn",
            "Extends naturally to diabetes (glucose twin) using the same engine"], 0.6, 5.35, 12, 1.8, 16)

# 15 ─ Thank you
s = slide()
text(s, "Thank you", 0.6, 2.0, 12, 1.2, 60, INK, bold=True)
text(s, "HridayTwin: predict the surge, explain it, simulate the fix, act before it happens.", 0.6, 3.3, 12, 0.8, 22, INK2)
text(s, ["Team [Team name]  |  [College / incubator]", "Repository: [GitHub link]   |   Demo video: [YouTube link]",
         "Built on MyZenn (myzenn.com)"], 0.6, 4.6, 12, 1.6, 16, SLATE)

out = DOCS / "HridayTwin_Presentation.pptx"
prs.save(out)
print("saved", out, "| example patient", er["id"])
